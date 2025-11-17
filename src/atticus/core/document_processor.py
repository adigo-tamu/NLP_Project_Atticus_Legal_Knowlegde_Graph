"""
Document processing pipeline for legal documents.

This module provides parsers for various document formats (PDF, DOCX, TXT)
and creates structured document representations.
"""

import hashlib
import uuid
from datetime import datetime
from pathlib import Path
from typing import List, Optional

import pdfplumber
from docx import Document as DocxDocument

from atticus.core.config import get_config
from atticus.core.exceptions import DocumentProcessingError
from atticus.core.logger import get_logger
from atticus.core.models import Document, DocumentChunk, DocumentType, Position, ProcessingStatus
from atticus.utils.text_utils import (
    chunk_text_hierarchical,
    clean_legal_text,
    count_tokens,
    extract_sections,
)

logger = get_logger(__name__)


class DocumentParser:
    """Base class for document parsers."""

    def __init__(self):
        """Initialize document parser."""
        self.config = get_config()

    def parse(self, file_path: str) -> str:
        """
        Parse document and extract text.

        Args:
            file_path: Path to document file

        Returns:
            Extracted text content
        """
        raise NotImplementedError("Subclasses must implement parse()")

    def extract_metadata(self, file_path: str) -> dict:
        """
        Extract metadata from document.

        Args:
            file_path: Path to document file

        Returns:
            Dictionary of metadata
        """
        return {
            "file_path": file_path,
            "file_name": Path(file_path).name,
            "file_size": Path(file_path).stat().st_size,
        }


class PDFParser(DocumentParser):
    """Parser for PDF documents."""

    def parse(self, file_path: str) -> str:
        """Parse PDF and extract text."""
        try:
            text_content = []

            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_content.append(text)

            full_text = "\n\n".join(text_content)
            cleaned_text = clean_legal_text(full_text)

            logger.info(f"Successfully parsed PDF: {file_path} ({len(text_content)} pages)")
            return cleaned_text

        except Exception as e:
            error_msg = f"Failed to parse PDF {file_path}: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg)

    def extract_metadata(self, file_path: str) -> dict:
        """Extract PDF metadata."""
        metadata = super().extract_metadata(file_path)

        try:
            with pdfplumber.open(file_path) as pdf:
                metadata.update(
                    {
                        "page_count": len(pdf.pages),
                        "pdf_metadata": pdf.metadata,
                    }
                )
        except Exception as e:
            logger.warning(f"Failed to extract PDF metadata: {e}")

        return metadata


class DOCXParser(DocumentParser):
    """Parser for DOCX documents."""

    def parse(self, file_path: str) -> str:
        """Parse DOCX and extract text."""
        try:
            doc = DocxDocument(file_path)

            paragraphs = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    paragraphs.append(paragraph.text.strip())

            full_text = "\n\n".join(paragraphs)
            cleaned_text = clean_legal_text(full_text)

            logger.info(f"Successfully parsed DOCX: {file_path} ({len(paragraphs)} paragraphs)")
            return cleaned_text

        except Exception as e:
            error_msg = f"Failed to parse DOCX {file_path}: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg)

    def extract_metadata(self, file_path: str) -> dict:
        """Extract DOCX metadata."""
        metadata = super().extract_metadata(file_path)

        try:
            doc = DocxDocument(file_path)
            core_props = doc.core_properties

            metadata.update(
                {
                    "author": core_props.author,
                    "title": core_props.title,
                    "created": core_props.created,
                    "modified": core_props.modified,
                    "paragraph_count": len(doc.paragraphs),
                }
            )
        except Exception as e:
            logger.warning(f"Failed to extract DOCX metadata: {e}")

        return metadata


class TXTParser(DocumentParser):
    """Parser for plain text documents."""

    def parse(self, file_path: str) -> str:
        """Parse TXT file and extract text."""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()

            cleaned_text = clean_legal_text(text)

            logger.info(f"Successfully parsed TXT: {file_path}")
            return cleaned_text

        except UnicodeDecodeError:
            # Try alternative encodings
            try:
                with open(file_path, "r", encoding="latin-1") as f:
                    text = f.read()
                cleaned_text = clean_legal_text(text)
                return cleaned_text
            except Exception as e:
                error_msg = f"Failed to parse TXT {file_path}: {str(e)}"
                logger.error(error_msg)
                raise DocumentProcessingError(error_msg)
        except Exception as e:
            error_msg = f"Failed to parse TXT {file_path}: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg)


class DocumentProcessor:
    """Main document processor for creating structured documents."""

    def __init__(self):
        """Initialize document processor."""
        self.config = get_config()
        self.parsers = {
            "pdf": PDFParser(),
            "docx": DOCXParser(),
            "doc": DOCXParser(),
            "txt": TXTParser(),
        }

    def process_file(self, file_path: str, document_type: Optional[DocumentType] = None) -> Document:
        """
        Process a legal document file.

        Args:
            file_path: Path to document file
            document_type: Optional document type classification

        Returns:
            Structured Document object
        """
        file_path = Path(file_path)

        if not file_path.exists():
            raise DocumentProcessingError(f"File not found: {file_path}")

        # Determine file format
        file_extension = file_path.suffix.lower().lstrip(".")

        if file_extension not in self.parsers:
            raise DocumentProcessingError(
                f"Unsupported file format: {file_extension}. "
                f"Supported formats: {list(self.parsers.keys())}"
            )

        # Parse document
        parser = self.parsers[file_extension]
        logger.info(f"Processing document: {file_path}")

        try:
            # Extract text
            text_content = parser.parse(str(file_path))

            # Extract metadata
            metadata = parser.extract_metadata(str(file_path))

            # Generate document ID
            doc_id = self._generate_document_id(str(file_path), text_content)

            # Detect document type if not provided
            if not document_type:
                document_type = self._detect_document_type(text_content)

            # Extract sections
            sections = extract_sections(text_content)
            section_titles = [s["title"] for s in sections if s.get("title")]

            # Create Document object
            document = Document(
                id=doc_id,
                title=metadata.get("title") or file_path.stem,
                content=text_content,
                document_type=document_type,
                file_path=str(file_path),
                file_format=file_extension,
                page_count=metadata.get("page_count"),
                sections=section_titles,
                status=ProcessingStatus.PENDING,
                metadata=metadata,
            )

            logger.info(
                f"Successfully processed document {doc_id}: "
                f"{len(text_content)} characters, "
                f"{count_tokens(text_content)} tokens, "
                f"{len(sections)} sections"
            )

            return document

        except Exception as e:
            error_msg = f"Error processing document {file_path}: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg)

    def create_chunks(self, document: Document) -> List[DocumentChunk]:
        """
        Create hierarchical chunks from document.

        Args:
            document: Document object

        Returns:
            List of DocumentChunk objects
        """
        logger.info(f"Creating chunks for document {document.id}")

        try:
            # Use hierarchical chunking
            chunk_dicts = chunk_text_hierarchical(
                document.content,
                section_size=self.config.processing.chunk_size * 2,
                clause_size=self.config.processing.chunk_size,
                overlap=self.config.processing.chunk_overlap,
            )

            chunks = []
            for chunk_dict in chunk_dicts:
                chunk_id = f"{document.id}_chunk_{chunk_dict['index']}"

                # Find position in document
                start_pos = document.content.find(chunk_dict["text"][:50])
                end_pos = start_pos + len(chunk_dict["text"]) if start_pos != -1 else 0

                chunk = DocumentChunk(
                    id=chunk_id,
                    document_id=document.id,
                    text=chunk_dict["text"],
                    position=Position(start=start_pos, end=end_pos),
                    chunk_index=chunk_dict["index"],
                    level=chunk_dict["level"],
                    parent_chunk_id=f"{document.id}_chunk_{chunk_dict['parent_index']}"
                    if chunk_dict.get("parent_index") is not None
                    else None,
                    token_count=count_tokens(chunk_dict["text"]),
                    metadata={"chunk_type": chunk_dict["chunk_type"]},
                )

                chunks.append(chunk)

            logger.info(f"Created {len(chunks)} chunks for document {document.id}")
            return chunks

        except Exception as e:
            error_msg = f"Error creating chunks for document {document.id}: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg, document_id=document.id)

    def _generate_document_id(self, file_path: str, content: str) -> str:
        """Generate unique document ID based on file path and content hash."""
        # Create hash of content for uniqueness
        content_hash = hashlib.sha256(content.encode()).hexdigest()[:8]

        # Use filename and hash
        file_name = Path(file_path).stem
        return f"doc_{file_name}_{content_hash}"

    def _detect_document_type(self, text: str) -> DocumentType:
        """
        Detect document type from content.

        Args:
            text: Document text

        Returns:
            Detected DocumentType
        """
        text_lower = text.lower()

        # Simple heuristic-based detection
        if any(keyword in text_lower for keyword in ["contract", "agreement", "parties hereby agree"]):
            if "service" in text_lower:
                return DocumentType.CONTRACT
            if "license" in text_lower:
                return DocumentType.LICENSE
            return DocumentType.AGREEMENT

        if "amendment" in text_lower:
            return DocumentType.AMENDMENT

        if "addendum" in text_lower:
            return DocumentType.ADDENDUM

        if "policy" in text_lower:
            return DocumentType.POLICY

        return DocumentType.OTHER

    def process_batch(self, file_paths: List[str]) -> List[Document]:
        """
        Process multiple documents in batch.

        Args:
            file_paths: List of file paths

        Returns:
            List of processed Documents
        """
        documents = []

        for file_path in file_paths:
            try:
                document = self.process_file(file_path)
                documents.append(document)
            except Exception as e:
                logger.error(f"Failed to process {file_path}: {e}")
                continue

        logger.info(f"Batch processing completed: {len(documents)}/{len(file_paths)} successful")
        return documents
